"""Identity-safe Baseball Card Data projections and marketing exports.

This service deliberately projects only factual standings from an AKS capture.
It never reads participant display names, and it keeps modelled scenario claims
out of the event projection.  The 24-column ``Card Data`` contract is stable
because it is consumed by an external marketing import workflow.
"""
from __future__ import annotations

import csv
import math
import re
from datetime import datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Iterable

from ..config import AksSettings
from ..integrations.aks.client import AksTimingClient
from ..integrations.aks.live_adapter import AksLiveIngestAdapter
from ..integrations.aks.recording import AksRecordingLoader
from .baseball_card_evidence import CardEvidenceProjection, configured_target, eligible_challengers


CARD_DATA_HEADERS = [
    "Card No", "Participant ID", "Participant Name", "Kart No", "Kart Display",
    "Laps When Max Passed", "Time Spent Ahead of Max", "Field Performance Rank", "Field Performance Display", "Start Position",
    "Peak Position", "Final Position", "Position Change", "Start Display", "Peak Display",
    "Final Display", "Change Display", "Primary Result Label", "Primary Result Statement", "Time Metric Label",
    "Field Metric Label", "Campaign", "Campaign Subtitle", "Photo Asset",
]

_REHEARSAL_SOURCE = Path(
    "/private/tmp/race-intelligence-aks-live-test/"
    "session-REHEARSAL_LIVE_CHECK-20260911T223558Z"
)
def _integer(value: Any) -> int | None:
    try:
        number = int(value)
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _recording_for(session_id: str) -> Path:
    """Resolve only the locally approved rehearsal capture for this phase.

    This is intentionally not a generic raw-capture browser: arbitrary local
    recordings must not become an implicit Event API data source.
    """
    if session_id in {"REHEARSAL_LIVE_CHECK", "REHEARSAL_AKS_4040"} and _REHEARSAL_SOURCE.is_dir():
        return _REHEARSAL_SOURCE
    raise FileNotFoundError("No approved local Baseball Card rehearsal recording is available for this session.")


def _event_observations(session_id: str) -> dict[str, dict[str, Any]]:
    """Project existing normalized events from already-durable offline evidence."""
    projection = CardEvidenceProjection()
    client = AksTimingClient(AksSettings(host=None))
    adapter = AksLiveIngestAdapter(session_id=session_id, client=client, publisher=projection,
                                   persist_raw=lambda line: None)
    # No new journal is needed: these frames are read from the approved durable
    # recording. No connection, broker, AIDP or ADB write is opened here.
    for message in AksRecordingLoader(_recording_for(session_id)).iter_messages():
        adapter.process_line(message.wire_line, recorded_at=message.recorded_at)
    return projection.result()


def _display_position(value: int | None) -> str | None:
    return f"P{value}" if value is not None else None


def _ordinal(value: int | None) -> str | None:
    if value is None:
        return None
    suffix = "th" if 10 <= value % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(value % 10, "th")
    return f"{value}{suffix}"


def _event_row(card_number: int, source: dict[str, Any]) -> dict[str, Any]:
    # Capture opened after the race had already progressed.  First observation
    # is retained in details, but is not presented as an actual race start.
    peak = source["peakPosition"]
    return {
        "Card No": card_number,
        "Participant ID": source["participantId"],
        "Participant Name": None,
        "Kart No": source["kartNo"],
        "Kart Display": f"KART {source['kartNo']}",
        "Laps When Max Passed": None, "Time Spent Ahead of Max": None, "Field Performance Rank": None, "Field Performance Display": None,
        "Start Position": source.get('startPosition'), "Peak Position": peak, "Final Position": None,
        "Position Change": None, "Start Display": _display_position(source.get('startPosition')), "Peak Display": _display_position(peak),
        "Final Display": None, "Change Display": None, "Primary Result Label": "FACTUAL EVENT",
        "Primary Result Statement": None, "Time Metric Label": None, "Field Metric Label": None,
        "Campaign": "MAX VS 100", "Campaign Subtitle": "RACE INTELLIGENCE", "Photo Asset": None,
        # Internal/API-only fields.  They never appear in Card Data.
        "status": "PARTIAL",
        "statusReasons": ([] if source.get('startPosition') is not None else ['NO_PROVEN_START_POSITION']) + ["NO_FINAL_CLASSIFICATION", "MISSING_CATCH_MODEL"],
        "evidenceDomain": "FACTUAL_OBSERVATION",
        "operationalContext": "REHEARSAL",
        "eventContractVersion": "1.0",
        "startPositionStatus": "GOVERNED_START_OBSERVED" if source.get('startPosition') is not None else "FIRST_OBSERVED",
        "firstObservedPosition": source["firstObservedPosition"],
        "latestObservedPosition": source["latestObservedPosition"],
        "latestObservedAt": source["latestObservedAt"],
        "observedLap": source.get('observedLap'),
        "gapToLeaderSeconds": source.get('gapToLeaderSeconds'),
        "intervalSeconds": source.get('intervalSeconds'),
        "latestCompletedLap": source.get('latestCompletedLap'),
        "bestLapSeconds": source.get('bestLapSeconds'),
        "bestLapNumber": source.get('bestLapNumber'),
        "metricProvenance": {
            'standings': 'AKS / timing.session.standings.overall.active',
            'laps': 'AKS / timing.analysis.laps / isValid=true / positive lap time',
            'bestLapObservedAt': source.get('bestLapObservedAt'),
            'finalStateAssessment': 'FINISHED_OBSERVED_NOT_OFFICIAL_CLASSIFICATION' if source.get('finishedStateObserved') else 'UNAVAILABLE',
            'captureResetBoundaryObserved': source.get('resetBoundaryObserved', False),
        },
    }


def _kart_sort_key(row: dict[str, Any]) -> tuple[int, str]:
    value = row["Kart No"]
    return (0, f"{value:08d}") if isinstance(value, int) else (1, str(value))


def build_serving_card_data(session_id: str, database) -> dict[str, Any]:
    """Post-session cutover through the existing Autonomous read-only integration."""
    records = database.fetch_all(
        "SELECT logical_session_id, source_epoch, participant_id, kart_number, "
        "latest_observed_position, peak_observed_position, observed_lap, gap_to_leader, "
        '"INTERVAL", best_valid_lap, start_position, final_position, position_change, '
        "updated_at, operational_context, evidence_domain, event_contract_version "
        "FROM V_EVENT_BASEBALL_CARD_EVIDENCE WHERE logical_session_id = :session_id",
        {"session_id": session_id},
    )
    if not records:
        raise FileNotFoundError("No published factual card evidence for this session/epoch.")
    epochs = {r["sourceEpoch"] for r in records}
    contexts = {r["operationalContext"] for r in records}
    if len(epochs) != 1 or len(contexts) != 1:
        raise ValueError("Published factual evidence mixes source epochs or contexts.")
    rows = []
    for r in records:
        if (r["logicalSessionId"] != session_id or r["evidenceDomain"] != "FACTUAL_OBSERVATION"
                or r["eventContractVersion"] != "1.0" or str(r["kartNumber"]).upper() == "SC"
                or r.get("finalPosition") is not None or r.get("positionChange") is not None):
            raise ValueError("Published factual evidence violates the serving contract.")
        kart = str(r["kartNumber"])
        source = {
            "participantId": r["participantId"], "kartNo": int(kart) if kart.isdigit() else kart,
            "peakPosition": r["peakObservedPosition"], "latestObservedPosition": r["latestObservedPosition"],
            "firstObservedPosition": None, "latestObservedAt": r["updatedAt"],
            "observedLap": r["observedLap"], "gapToLeaderSeconds": r["gapToLeader"],
            "intervalSeconds": r["interval"], "bestLapSeconds": r["bestValidLap"],
            "startPosition": r["startPosition"],
        }
        row = _event_row(0, source)
        row["operationalContext"] = r["operationalContext"]
        rows.append(row)
    rows.sort(key=_kart_sort_key)
    for number, row in enumerate(rows, 1):
        row["Card No"] = number
    return {"sessionId": session_id, "mode": "event", "rows": rows, "metadata": {
        "sourceType": "AIDP_GOLD_EVENT", "sourceMode": "REPLAY", "sourceEpoch": next(iter(epochs)),
        "sourceState": "READY", "operationalContext": next(iter(contexts)), "rowCount": len(rows),
        "eventContractVersion": "1.0", "evidenceDomain": "FACTUAL_OBSERVATION",
        "finalClassificationSupport": "NOT_SUPPORTED", "modelledMetricsSupport": "NOT_SUPPORTED_FOR_EVENT_MODE",
    }}


def build_event_card_data(session_id: str) -> dict[str, Any]:
    target = configured_target()
    rows = [_event_row(0, source) for source in _event_observations(session_id).values()]
    rows.sort(key=_kart_sort_key)
    for card_number, row in enumerate(rows, start=1):
        row["Card No"] = card_number
    return {
        "sessionId": session_id,
        "mode": "event",
        "rows": rows,
        "metadata": {
            "sourceType": "LOCAL_APPROVED_AKS_REHEARSAL_RECORDING",
            "operationalContext": "REHEARSAL",
            "eventContractVersion": "1.0",
            "evidenceDomain": "FACTUAL_OBSERVATION",
            "rowCount": len(rows),
            "modelledMetricsSupport": "NOT_SUPPORTED_FOR_EVENT_MODE",
            "finalClassificationSupport": "NOT_SUPPORTED",
            "targetKartNumber": target,
            "eligibleChallengerCount": len(eligible_challengers([str(row['Kart No']) for row in rows], target)),
            "timingCatchSupport": "NOT_SUPPORTED: no approved common-clock catch threshold; capture begins mid-race",
        },
    }


def build_modelled_card_data(session_id: str) -> dict[str, Any]:
    """An explicit safe placeholder until an existing scenario result is supplied.

    The Event endpoint has no scenario id/request result to project, therefore
    calculating cards here would be a second simulation engine.  The UI makes
    this limitation visible rather than inventing modelled values.
    """
    return {
        "sessionId": session_id, "mode": "modelled", "rows": [],
        "metadata": {
            "sourceType": "NO_SCENARIO_RESULT_BOUND_TO_EVENT_SESSION",
            "evidenceDomain": "MODELLED_SCENARIO",
            "rowCount": 0,
            "support": "NOT_SUPPORTED",
            "reason": "Select a completed Synthetic 100 scenario before modelled card projection is enabled.",
        },
    }


def card_data(session_id: str, mode: str) -> dict[str, Any]:
    if mode == "event":
        return build_event_card_data(session_id)
    if mode == "modelled":
        return build_modelled_card_data(session_id)
    raise ValueError("Unsupported Baseball Card Data mode.")


def modelled_export_payload(payload: dict[str, Any]) -> dict[str, Any]:
    """Serialize the existing frontend participant projection; never simulate.

    This transient export is client-supplied MODELLED data, never persisted or
    promoted into the factual Event domain. Strip arbitrary strings and names.
    """
    rows = payload.get('rows')
    if payload.get('mode') != 'modelled' or not isinstance(rows, list) or len(rows) != 100:
        raise ValueError('Exactly 100 MODELLED participant rows are required.')
    safe, ids, karts = [], set(), set()
    for source in rows:
        identifier, kart = source.get('Participant ID'), source.get('Kart No')
        if not isinstance(identifier, str) or not re.fullmatch(r'SYN_[0-9]{3}', identifier):
            raise ValueError('Only anonymous synthetic participant identifiers are accepted.')
        if type(kart) is not int or not 1 <= kart <= 100 or identifier in ids or kart in karts:
            raise ValueError('Unique synthetic kart keys 1–100 are required.')
        if source.get('evidenceDomain') != 'MODELLED_SCENARIO':
            raise ValueError('Factual evidence cannot enter a modelled export.')
        ids.add(identifier); karts.add(kart)
        row = dict.fromkeys(CARD_DATA_HEADERS)
        for key in ['Laps When Max Passed', 'Time Spent Ahead of Max', 'Field Performance Rank', 'Start Position', 'Peak Position', 'Final Position', 'Position Change']:
            value = source.get(key)
            if value is not None and (type(value) not in (int, float) or not math.isfinite(value)):
                raise ValueError('Modelled metrics must be finite numbers or null.')
            if value is not None and key != 'Position Change' and value < 0:
                raise ValueError('Modelled duration and rank cannot be negative.')
            if value is not None and key in {'Start Position', 'Peak Position', 'Final Position'} and (not 1 <= value <= 101 or int(value) != value):
                raise ValueError('Modelled rank must be P1–P101.')
            if value is not None and key == 'Field Performance Rank' and (not 1 <= value <= 100 or int(value) != value):
                raise ValueError('Field Performance Rank must be an integer from 1 to 100.')
            row[key] = value
        row.update({'Card No': len(safe) + 1, 'Participant ID': identifier, 'Kart No': kart, 'Kart Display': f'KART {kart}',
                    'Campaign': 'MAX VS 100', 'Campaign Subtitle': 'RACE INTELLIGENCE',
                    'Primary Result Label': 'PRIMARY RESULT', 'Primary Result Statement': 'WHEN MAX PASSED',
                    'Time Metric Label': 'TIME SPENT AHEAD OF MAX', 'Field Metric Label': 'FIELD PERFORMANCE'})
        row['Field Performance Display'] = _ordinal(row['Field Performance Rank'])
        for label in ('Start', 'Peak', 'Final'):
            row[f'{label} Display'] = _display_position(row[f'{label} Position'])
        change = row['Position Change']
        row['Change Display'] = None if change is None else f'{change:+g}' if change > 0 else f'{change:g}'
        safe.append(row)
    return {'sessionId': 'Synthetic100', 'mode': 'modelled', 'rows': safe,
            'metadata': {'evidenceDomain': 'MODELLED_SCENARIO', 'rowCount': 100,
                         'sourceType': 'CLIENT_EXISTING_SCENARIO_CARD_PROJECTION', 'persisted': False}}


def card_data_values(row: dict[str, Any]) -> list[Any]:
    """The only projection allowed into Toni's fixed workbook contract."""
    values = [row.get(header) for header in CARD_DATA_HEADERS]
    # Internal/API time is seconds. XLSX stores durations as fractions
    # of a day, so the fixed Card Data sheet remains a real Excel duration.
    time_ahead_index = CARD_DATA_HEADERS.index("Time Spent Ahead of Max")
    if values[time_ahead_index] is not None:
        values[time_ahead_index] = float(values[time_ahead_index]) / 86_400
    return values


def _add_sheet(workbook: Any, name: str, headers: list[str], rows: Iterable[list[Any]]) -> Any:
    from openpyxl.styles import Font, PatternFill
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="071B34")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 54)
    return sheet


def build_card_data_workbook(payload: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    card = _add_sheet(workbook, "Card Data", CARD_DATA_HEADERS, (card_data_values(row) for row in payload["rows"]))
    for row in range(2, card.max_row + 1):
        card.cell(row, 6).number_format = "0.0"
        card.cell(row, 7).number_format = "[m]:ss.0"
        card.cell(row, 8).number_format = "0"
    metadata = payload["metadata"]
    _add_sheet(workbook, "Metadata", ["Key", "Value"], [
        ["session_id", payload["sessionId"]], ["generated_at", metadata.get("generatedAt") or datetime.now(timezone.utc).isoformat()],
        ["source_mode", metadata.get("sourceMode")], ["state_version", metadata.get("stateVersion")],
        ["card_data_mode", payload["mode"]], ["row_count", metadata.get("rowCount")],
        ["operational_context", metadata.get("operationalContext")],
        ["event_contract_version", metadata.get("eventContractVersion")], ["source", metadata.get("sourceType")],
        ["identity_boundary", "Participant Name is intentionally blank. Authorised content teams enrich externally using Kart No."],
        ["methodology_boundary", "MODELLED_SCENARIO and FACTUAL_OBSERVATION are separate; this export contains only its selected mode."],
        ["modelled_methodology_note", '"Laps When Max Passed" uses the existing modelled timing-catch calculation. It is not an observed physical overtake.'],
    ])
    _add_sheet(workbook, "Data Dictionary", ["Column", "Meaning"], [
        ["Participant Name", "Always blank. Race Intelligence does not receive or export driver identity for this workflow."],
        ["Kart No", "External join key for authorised identity/photo enrichment outside Race Intelligence."],
        ["Laps When Max Passed", "Existing modelled timing-catch race progress; blank without approved scenario support."],
        ["Time Spent Ahead of Max", "Existing modelled timing-catch clock value; blank for factual event data."],
        ["Field Performance Rank", "1 plus the number of comparable participants with a strictly earlier timing catch."],
        ["Start Position", "Factual start only when capture is proven to cover race start."],
        ["Peak Position", "Best factual observed position (lowest numeric position) in captured evidence."],
        ["Final Position", "Blank unless an approved factual final-classification rule exists."],
        ["Position Change", "Start Position minus Final Position; blank without final classification."],
    ])
    output = BytesIO(); workbook.save(output)
    return output.getvalue()


def build_card_data_csv(payload: dict[str, Any]) -> str:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(CARD_DATA_HEADERS)
    duration = CARD_DATA_HEADERS.index("Time Spent Ahead of Max")
    for row in payload["rows"]:
        values = [row.get(header) for header in CARD_DATA_HEADERS]
        if values[duration] is not None:
            tenths = round(float(values[duration]) * 10)
            values[duration] = f"{tenths // 600:02d}:{(tenths % 600) // 10:02d}.{tenths % 10}"
        writer.writerow(values)
    return output.getvalue()
