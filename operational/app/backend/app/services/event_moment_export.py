"""Workbook exporter for the identity-safe Event Moments handoff."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any


LEGAL_NOTE = (
    "This workbook contains race/event data only. Driver names, photographs and media assets are intentionally "
    "excluded and must be joined externally by the authorised content team."
)

# Versioned additive LIVE export. The established multi-sheet serving export
# above/below retains its original contract without renaming existing columns.
FROZEN_COLUMNS = [
    ('Moment No',None),('Moment ID','moment_id'),('Session ID','session_id'),('Captured At','captured_at'),('Observed At','observed_at'),
    ('Operational Context','operational_context'),('Evidence Domain','evidence_domain'),('Moment Type','moment_type'),
    ('Capture Origin','capture_origin'),('Operator Asserted','operator_asserted'),('Primary Participant ID','primary_participant_id'),
    ('Primary Kart No','primary_kart_number'),('Secondary Participant ID','secondary_participant_id'),('Secondary Kart No','secondary_kart_number'),
    ('Observed Leader Kart','observed_leader_kart'),('Observed Field Size','observed_field_size'),('Observed Lap','observed_lap'),
    ('Previous Position','previous_position'),('Current Position','latest_observed_position'),('Previous Gap','previous_gap'),
    ('Current Gap','gap_to_leader'),('Gap Delta','gap_delta'),('Peak Observed Position','peak_observed_position'),('Interval','interval'),
    ('Best Valid Lap','best_valid_lap'),('Headline','headline'),('Operator Note','operator_note'),('Review Status','review_status'),
    ('Social Selected','social_selected'),('Disclosure','disclosure')]


def frozen_moments_export(snapshot, format):
    """One detached collection; never consult current race state during export."""
    rows=[[i if key is None else row.get(key) for _,key in FROZEN_COLUMNS] for i,row in enumerate(snapshot['moments'],1)]
    headers=[h for h,_ in FROZEN_COLUMNS]
    if format=='csv':
        import csv
        from io import StringIO
        out=StringIO(newline=''); writer=csv.writer(out);writer.writerow(headers)
        for row in rows:
            writer.writerow(["'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for v in row])
        return out.getvalue().encode('utf-8-sig')
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    book=Workbook(); sheet=book.active;sheet.title='Event Moments';sheet.append(headers)
    for row in rows: sheet.append(row)
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value,str): cell.data_type='s'
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='071B34')
    for i in range(1,len(headers)+1):sheet.column_dimensions[get_column_letter(i)].width=min(44,max(20,len(headers[i-1])+2))
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    metadata=book.create_sheet('Metadata');metadata.append(['Key','Value'])
    for key,value in snapshot['metadata'].items():metadata.append([key,value])
    metadata.column_dimensions['A'].width=30;metadata.column_dimensions['B'].width=48
    out=BytesIO();book.save(out);return out.getvalue()


def _iso(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat().replace("+00:00", "Z")
    return value


def build_social_data_workbook(*, session_id: str, moments: list[dict[str, Any]], snapshots: list[dict[str, Any]]) -> bytes:
    """Build a flat, external-join-safe XLSX handoff in memory.

    openpyxl is a deliberate backend dependency rather than a browser export so
    filters and sheet definitions share the same service result as the UI.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    by_moment = {moment["moment_id"]: moment for moment in moments}
    workbook = Workbook()
    workbook.remove(workbook.active)

    def add_sheet(name: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="071B34")
        for row in rows:
            sheet.append([_iso(row.get(header)) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 12), 48)

    card_headers = ["SESSION_ID", "MOMENT_ID", "KART_NUMBER", "PARTICIPANT_ID", "OBSERVED_AT_UTC", "START_POSITION", "PEAK_POSITION", "CURRENT_POSITION", "POSITION_CHANGE", "LAP_NUMBER", "GAP_TO_LEADER_SECONDS", "STATUS", "EVIDENCE_DOMAIN", "OPERATIONAL_CONTEXT", "EVENT_CONTRACT_VERSION"]
    cards = [{
        "SESSION_ID": snapshot.get("session_id"), "MOMENT_ID": snapshot.get("moment_id"),
        "KART_NUMBER": snapshot.get("external_participant_number"), "PARTICIPANT_ID": snapshot.get("participant_id"),
        "OBSERVED_AT_UTC": snapshot.get("observed_at"), "START_POSITION": snapshot.get("start_position"),
        "PEAK_POSITION": snapshot.get("peak_position"), "CURRENT_POSITION": snapshot.get("position"),
        "POSITION_CHANGE": snapshot.get("position_change"), "LAP_NUMBER": snapshot.get("lap_number"),
        "GAP_TO_LEADER_SECONDS": snapshot.get("gap_first_seconds"), "STATUS": snapshot.get("status"),
        "EVIDENCE_DOMAIN": snapshot.get("evidence_domain"), "OPERATIONAL_CONTEXT": snapshot.get("operational_context"),
        "EVENT_CONTRACT_VERSION": snapshot.get("event_contract_version"),
    } for snapshot in snapshots]
    add_sheet("BASEBALL_CARD_DATA", card_headers, cards)

    moment_headers = ["MOMENT_ID", "SESSION_ID", "MOMENT_TYPE", "TITLE", "DESCRIPTION", "TRIGGER_MODE", "TRIGGER_RULE", "OBSERVED_AT_UTC", "PRIMARY_PARTICIPANT_ID", "KART_NUMBER", "RELATED_PARTICIPANT_ID", "RELATED_KART_NUMBER", "PRIOR_VALUE", "CURRENT_VALUE", "DELTA_VALUE", "UNIT", "IMPORTANCE", "SOCIAL_PRIORITY", "OPERATOR_NOTE", "EVIDENCE_DOMAIN", "SOURCE_SNAPSHOT_ID"]
    event_rows = [{
        "MOMENT_ID": row.get("moment_id"), "SESSION_ID": row.get("session_id"), "MOMENT_TYPE": row.get("moment_type"),
        "TITLE": row.get("title"), "DESCRIPTION": row.get("description"), "TRIGGER_MODE": row.get("trigger_mode"),
        "TRIGGER_RULE": row.get("trigger_rule"), "OBSERVED_AT_UTC": row.get("observed_at"),
        "PRIMARY_PARTICIPANT_ID": row.get("primary_participant_id"), "KART_NUMBER": row.get("external_participant_number"),
        "RELATED_PARTICIPANT_ID": row.get("related_participant_id"), "RELATED_KART_NUMBER": row.get("related_external_participant_number"),
        "PRIOR_VALUE": row.get("prior_value"), "CURRENT_VALUE": row.get("current_value"), "DELTA_VALUE": row.get("delta_value"),
        "UNIT": row.get("unit"), "IMPORTANCE": row.get("importance"), "SOCIAL_PRIORITY": row.get("social_priority"),
        "OPERATOR_NOTE": row.get("operator_note"), "EVIDENCE_DOMAIN": row.get("evidence_domain"), "SOURCE_SNAPSHOT_ID": row.get("source_snapshot_id"),
    } for row in moments]
    add_sheet("EVENT_MOMENTS", moment_headers, event_rows)
    snapshot_headers = ["SNAPSHOT_ID", "MOMENT_ID", "SESSION_ID", "PARTICIPANT_ID", "KART_NUMBER", "OBSERVED_AT_UTC", "POSITION", "LAP_NUMBER", "GAP_TO_LEADER_SECONDS", "GAP_TO_PREVIOUS_SECONDS", "POSITION_CHANGE", "AT_LOOP", "STATUS", "START_POSITION", "PEAK_POSITION", "EVIDENCE_DOMAIN"]
    add_sheet("PARTICIPANT_SNAPSHOTS", snapshot_headers, [{
        "SNAPSHOT_ID": row.get("snapshot_id"), "MOMENT_ID": row.get("moment_id"), "SESSION_ID": row.get("session_id"),
        "PARTICIPANT_ID": row.get("participant_id"), "KART_NUMBER": row.get("external_participant_number"),
        "OBSERVED_AT_UTC": row.get("observed_at"), "POSITION": row.get("position"), "LAP_NUMBER": row.get("lap_number"),
        "GAP_TO_LEADER_SECONDS": row.get("gap_first_seconds"), "GAP_TO_PREVIOUS_SECONDS": row.get("gap_previous_seconds"),
        "POSITION_CHANGE": row.get("position_change"), "AT_LOOP": row.get("at_loop"), "STATUS": row.get("status"),
        "START_POSITION": row.get("start_position"), "PEAK_POSITION": row.get("peak_position"), "EVIDENCE_DOMAIN": row.get("evidence_domain"),
    } for row in snapshots])
    for name, selected in {
        "PREDICTION_CHANGES": {"PREDICTION_CHANGE", "TARGET_CHANGE"}, "POSITION_GAP_CHANGES": {"POSITION_CHANGE", "GAP_CHANGE", "LEADER_CHANGE"},
        "CONFIRMED_PASSES": {"CONFIRMED_PASS"}, "FASTEST_LAPS": {"FASTEST_LAP"},
        "OPERATOR_MARKERS": {"PRE_RACE", "START", "MID_RACE", "KEY_MOMENT", "CONTROL_ROOM_REACTION", "END_OF_CAPTURE", "FINISH", "MANUAL_MARKER"},
    }.items():
        add_sheet(name, moment_headers, [row for row in event_rows if row["MOMENT_TYPE"] in selected])
    add_sheet("DATA_DICTIONARY", ["FIELD", "MEANING"], [
        {"FIELD": "KART_NUMBER", "MEANING": "External participant number. Join to authorised creative identity assets outside Race Intelligence."},
        {"FIELD": "PARTICIPANT_ID", "MEANING": "Canonical operational participant identifier. It is not a personal identity."},
        {"FIELD": "CONFIRMED_PASS", "MEANING": "Explicit factual event or operator confirmation only. Never inferred from standings, timing, or GPS."},
        {"FIELD": "EVIDENCE_DOMAIN", "MEANING": "FACTUAL_OBSERVATION and MODELLED_SCENARIO remain separate."},
    ])
    add_sheet("METADATA", ["KEY", "VALUE"], [
        {"KEY": "SESSION_ID", "VALUE": session_id}, {"KEY": "LEGAL_NOTE", "VALUE": LEGAL_NOTE},
        {"KEY": "CAVEAT", "VALUE": "Observed standings are not final classification without explicit factual support."},
        {"KEY": "CAVEAT", "VALUE": "Model predictions are not facts."},
        {"KEY": "CAVEAT", "VALUE": "Physical passes are only factual or operator confirmed."},
    ])
    output = BytesIO(); workbook.save(output)
    return output.getvalue()
